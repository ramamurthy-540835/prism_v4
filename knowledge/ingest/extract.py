from __future__ import annotations

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass
from html import unescape
from pathlib import Path
from typing import Iterable

MAX_BYTES = 50 * 1024 * 1024
MAX_PAGES = 500


class ExtractionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedBlock:
    text: str
    page: int
    section_path: str = ""
    is_table: bool = False


@dataclass(frozen=True)
class ExtractedDocument:
    blocks: list[ExtractedBlock]
    page_count: int


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", unescape(text))).strip()


def _heading(text: str) -> str | None:
    first = text.strip().splitlines()[0] if text.strip() else ""
    if re.match(r"^\d+(?:\.\d+)*\s+\S+", first) or re.match(r"^(abstract|summary|references|appendix)\b", first, re.I):
        return first[:300]
    return None


def _with_sections(blocks: Iterable[ExtractedBlock]) -> list[ExtractedBlock]:
    path: list[str] = []
    result: list[ExtractedBlock] = []
    for block in blocks:
        heading = _heading(block.text)
        if heading:
            if heading.lower().startswith("references"):
                path = ["References"]
            else:
                level = heading.split()[0].count(".") + 1 if heading[0].isdigit() else 1
                path = path[: level - 1] + [heading]
        result.append(ExtractedBlock(block.text, block.page, " > ".join(path), block.is_table))
    return result


def _strip_repeating_margins(pages: list[str]) -> list[str]:
    if len(pages) < 3:
        return pages
    candidates: Counter[str] = Counter()
    split_pages = [page.splitlines() for page in pages]
    for lines in split_pages:
        for line in lines[:2] + lines[-2:]:
            normal = _clean(line).lower()
            if len(normal) >= 12:
                candidates[normal] += 1
    repeated = {line for line, n in candidates.items() if n / len(pages) > 0.5}
    return ["\n".join(line for line in lines if _clean(line).lower() not in repeated) for lines in split_pages]


def _pdf(data: bytes) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception:
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages = [(page.extract_text() or "") for page in pdf.pages]
        except Exception as exc:
            raise ExtractionError(f"PDF text extraction failed: {exc}") from exc
    if len(pages) > MAX_PAGES:
        raise ExtractionError("file exceeds 500 pages")
    pages = _strip_repeating_margins(pages)
    blocks = _with_sections(ExtractedBlock(_clean(page), i + 1) for i, page in enumerate(pages) if _clean(page))
    if not blocks:
        raise ExtractionError("PDF contains no extractable text")
    return ExtractedDocument(blocks, len(pages))


def _docx(data: bytes) -> ExtractedDocument:
    from docx import Document
    doc = Document(io.BytesIO(data))
    blocks = [ExtractedBlock(_clean(p.text), 1) for p in doc.paragraphs if _clean(p.text)]
    for table in doc.tables:
        rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        if rows:
            blocks.append(ExtractedBlock("\n".join("| " + " | ".join(row) + " |" for row in rows), 1, is_table=True))
    return ExtractedDocument(_with_sections(blocks), 1)


def _tabular(data: bytes, suffix: str) -> ExtractedDocument:
    rows: list[list[str]] = []
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8", errors="replace"))))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in wb.worksheets:
            rows.append([f"## {sheet.title}"])
            rows.extend([["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)])
    text = "\n".join("| " + " | ".join(row) + " |" for row in rows if row)
    return ExtractedDocument(_with_sections([ExtractedBlock(text, 1, is_table=True)]), 1)


def extract_bytes(data: bytes, filename: str, mime_type: str) -> ExtractedDocument:
    if len(data) > MAX_BYTES:
        raise ExtractionError("file exceeds 50 MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf" or mime_type == "application/pdf":
        return _pdf(data)
    if suffix == ".docx":
        return _docx(data)
    if suffix in {".csv", ".xlsx"}:
        return _tabular(data, suffix)
    if suffix in {".html", ".htm"}:
        try:
            from bs4 import BeautifulSoup
            text = BeautifulSoup(data, "html.parser").get_text("\n")
        except ImportError:
            text = re.sub(r"<[^>]+>", " ", data.decode("utf-8", errors="replace"))
    else:
        text = data.decode("utf-8", errors="replace")
    text = _clean(text)
    if not text:
        raise ExtractionError("no extractable text")
    return ExtractedDocument(_with_sections([ExtractedBlock(text, 1)]), 1)


def extract_with_document_ai(data: bytes, processor_name: str, location: str) -> ExtractedDocument:
    """OCR fallback for scanned PDFs; caller must provide a configured processor."""
    from google.cloud import documentai
    client = documentai.DocumentProcessorServiceClient(
        client_options={"api_endpoint": f"{location}-documentai.googleapis.com"}
    )
    result = client.process_document(request={"name": processor_name, "raw_document": {"content": data, "mime_type": "application/pdf"}}).document
    pages: dict[int, list[str]] = {}
    for page_number, page in enumerate(result.pages, start=1):
        parts: list[str] = []
        for layout in [p.layout for p in page.paragraphs]:
            text = "".join(result.text[s.start_index or 0:s.end_index] for s in layout.text_anchor.text_segments)
            if _clean(text):
                parts.append(_clean(text))
        pages[page_number] = parts
    if len(pages) > MAX_PAGES:
        raise ExtractionError("file exceeds 500 pages")
    return ExtractedDocument(_with_sections(ExtractedBlock("\n".join(parts), page) for page, parts in pages.items() if parts), len(pages))
